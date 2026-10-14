from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase as DjangoTestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from accounts.models import User
from projects.models import Module, Project
from .models import ImportBatch, TestCase

H = ['MODULE', 'SUB-MODULE', 'TEST CASE', 'STEPS', 'EXPECTED RESULT', 'BACKEND/FRONTEND', 'PRIORITY']


def xlsx(rows, headers=H, name='cases.xlsx'):
    wb = Workbook(); ws = wb.active; ws.title = 'Test cases'
    ws.append(headers)
    for r in rows:
        ws.append(r)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.data_type = 's'           # store as text, like typing '=1+1 with a leading apostrophe
    out = BytesIO(); wb.save(out)
    return SimpleUploadedFile(name, out.getvalue())


class ImportTests(DjangoTestCase):
    def setUp(self):
        self.qa = User.objects.create_user('qa1', password='x', role='qa')
        self.dev = User.objects.create_user('dev1', password='x', role='developer')
        self.p = Project.objects.create(name='Checkout'); self.p.members.add(self.qa, self.dev)
        self.mod = Module.objects.create(project=self.p, name='Payments')
        self.existing = TestCase.objects.create(
            project=self.p, module=self.mod, title='Pay by card', steps='old steps', expected_result='old exp',
            priority='high', created_by=self.qa)
        self.client.force_login(self.qa)

    def u(self, name, *args):
        return reverse(name, args=[self.p.pk, *args])

    def upload(self, file):
        return self.client.post(self.u('import_home'), {'file': file}, follow=True)

    def batch(self):
        return ImportBatch.objects.get()

    GOOD = ['Cart', 'Coupons', 'Apply coupon', '1. Add item', 'Price drops', 'Backend', 'High']

    def test_template_has_headers_and_dropdowns(self):
        r = self.client.get(self.u('import_template'))
        wb = load_workbook(BytesIO(r.content))
        ws = wb['Test cases']
        self.assertEqual([c.value for c in ws[1]], H)
        self.assertEqual(len(ws.data_validations.dataValidation), 2)
        self.assertIn('How to fill', wb.sheetnames)

    def test_preview_classifies_rows_and_saves_nothing(self):
        rows = [self.GOOD,
                ['Payments', '', 'Pay by card', 's', 'e', '', ''],           # already exists
                ['Cart', '', '', 's', 'e', '', ''],                           # missing title
                ['Cart', '', 'Bad layer', 's', 'e', 'Middle', '']]           # invalid layer
        self.upload(xlsx(rows))
        b = self.batch()
        self.assertEqual([r['status'] for r in b.rows], ['new', 'duplicate', 'invalid', 'invalid'])
        self.assertEqual(TestCase.objects.count(), 1)                        # nothing saved yet
        self.assertEqual(Module.objects.count(), 1)

    def test_confirm_skip_creates_new_only_and_logs_counts(self):
        self.upload(xlsx([self.GOOD, ['Payments', '', 'Pay by card', 'NEW', 'NEW', '', ''],
                          ['Cart', '', '', 's', 'e', '', '']]))
        self.client.post(self.u('import_confirm', self.batch().pk), {'policy': 'skip'})
        b = self.batch()
        self.assertEqual((b.created_count, b.updated_count, b.skipped_count, b.rejected_count), (1, 0, 1, 1))
        self.assertEqual(b.status, 'done')
        self.assertTrue(TestCase.objects.filter(title='Apply coupon', module__name='Cart', sub_module__name='Coupons').exists())
        self.existing.refresh_from_db(); self.assertEqual(self.existing.steps, 'old steps')

    def test_confirm_update_changes_steps_but_keeps_blank_priority(self):
        self.upload(xlsx([['payments', '', 'PAY BY CARD', 'NEW steps', 'NEW exp', '', '']]))   # different case, still a match
        self.client.post(self.u('import_confirm', self.batch().pk), {'policy': 'update'})
        self.existing.refresh_from_db()
        self.assertEqual((self.existing.steps, self.existing.priority), ('NEW steps', 'high'))
        self.assertEqual(TestCase.objects.count(), 1)

    def test_duplicate_inside_file_is_rejected(self):
        self.upload(xlsx([self.GOOD, self.GOOD]))
        self.assertEqual([r['status'] for r in self.batch().rows], ['new', 'invalid'])

    def test_double_confirm_does_not_import_twice(self):
        self.upload(xlsx([self.GOOD]))
        url = self.u('import_confirm', self.batch().pk)
        self.client.post(url); self.client.post(url)
        self.assertEqual(TestCase.objects.filter(title='Apply coupon').count(), 1)

    def test_bad_files_are_refused_with_a_message(self):
        r = self.upload(SimpleUploadedFile('notes.txt', b'hello'))
        self.assertContains(r, 'Only .xlsx')
        r = self.upload(SimpleUploadedFile('fake.xlsx', b'not a zip'))
        self.assertContains(r, 'not a valid .xlsx')
        r = self.upload(xlsx([self.GOOD], headers=['MODULE', 'TEST CASE']))
        self.assertContains(r, 'Missing column')
        self.assertEqual(ImportBatch.objects.count(), 0)

    def test_header_aliases_allow_reimport_of_our_csv_layout(self):
        h = ['module', 'Sub Module', 'Test case', 'Steps to reproduce', 'Expected result', 'Backend/Frontend', 'priority']
        self.upload(xlsx([self.GOOD], headers=h))
        self.assertEqual(self.batch().rows[0]['status'], 'new')

    def test_error_report_contains_rejected_rows_and_text_stays_text(self):
        self.upload(xlsx([self.GOOD, ['Cart', '', '=1+1', '', 'e', 'Middle', '']]))
        r = self.client.get(self.u('import_errors', self.batch().pk))
        ws = load_workbook(BytesIO(r.content))['Test cases']
        self.assertEqual([c.value for c in ws[1]][-1], 'ERROR')
        self.assertEqual(ws.max_row, 2)                       # header + the one rejected row
        self.assertEqual(ws['C2'].value, '=1+1'); self.assertEqual(ws['C2'].data_type, 's')
        self.assertIn('Steps is required', ws['H2'].value); self.assertEqual(ws['F2'].value, 'Middle')

    def test_error_report_roundtrip_reimports_after_fixing(self):
        self.upload(xlsx([['Cart', '', 'T', '', 'e', '', '']]))
        report = load_workbook(BytesIO(self.client.get(self.u('import_errors', self.batch().pk)).content))
        report['Test cases']['D2'] = 'fixed steps'
        out = BytesIO(); report.save(out)
        self.upload(SimpleUploadedFile('fixed.xlsx', out.getvalue()))
        self.assertEqual(ImportBatch.objects.get(status='preview').rows[0]['status'], 'new')

    def test_roles_and_isolation(self):
        self.client.force_login(self.dev)
        self.assertEqual(self.client.get(self.u('import_home')).status_code, 403)
        self.client.force_login(self.qa)
        self.upload(xlsx([self.GOOD])); pk = self.batch().pk
        other = User.objects.create_user('qa2', password='x', role='qa'); self.p.members.add(other)
        self.client.force_login(other)                         # another member cannot see someone's preview
        self.assertEqual(self.client.get(self.u('import_detail', pk)).status_code, 404)
        self.assertEqual(self.client.post(self.u('import_confirm', pk)).status_code, 404)

    def test_csv_export_neutralises_formulas(self):
        TestCase.objects.create(project=self.p, module=self.mod, title='=HYPERLINK("http://x")', steps='+1',
                                expected_result='ok', created_by=self.qa)
        body = self.client.get(self.u('cases_export')).content.decode()
        self.assertIn("'=HYPERLINK", body); self.assertIn("'+1", body)
